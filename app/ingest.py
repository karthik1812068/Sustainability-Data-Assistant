"""Read PDFs and Excel into text chunks."""

from __future__ import annotations

import re
from pathlib import Path

import pymupdf
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def ingest_file(path: str | Path, filename: str) -> dict:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        return ingest_xlsx(path, filename)
    if suffix == ".pdf":
        return ingest_pdf(path, filename)
    text = path.read_text(encoding="utf-8", errors="replace")
    return {
        "kind": "text",
        "chunks": [{"loc_type": "page", "locator": "file", "text": text, "ordinal": 0}],
        "preview": text[:2000],
    }


def ingest_pdf(path: Path, filename: str) -> dict:
    doc = pymupdf.open(path)
    chunks = []
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text("text") or ""
        text = _clean_pdf_text(text)
        pages.append({"page": i + 1, "text": text, "image_only": len(text.strip()) < 40})
        for ordinal, block in enumerate(_split_blocks(text)):
            chunks.append(
                {
                    "loc_type": "page",
                    "locator": f"page {i + 1}",
                    "text": block,
                    "ordinal": i * 100 + ordinal,
                }
            )
        if not text.strip():
            chunks.append(
                {
                    "loc_type": "page",
                    "locator": f"page {i + 1}",
                    "text": f"[No extractable text on page {i + 1} of {filename}. The page may be a scanned image.]",
                    "ordinal": i * 100,
                }
            )
    return {"kind": "pdf", "page_count": doc.page_count, "pages": pages, "chunks": chunks, "preview": "\n".join(p["text"] for p in pages)[:4000]}


def ingest_xlsx(path: Path, filename: str) -> dict:
    wb = load_workbook(path, data_only=True)
    chunks = []
    sheets = []
    ordinal = 0
    for ws in wb.worksheets:
        rows = []
        used = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, max_col=ws.max_column or 1):
            values = []
            coords = []
            for cell in row:
                if cell.value is None or str(cell.value).strip() == "":
                    continue
                values.append(str(cell.value).strip())
                coords.append(cell.coordinate)
                used.append(cell.coordinate)
            if values:
                rows.append(" | ".join(values))
                first, last = coords[0], coords[-1]
                locator = f"{ws.title}!{first}:{last}" if first != last else f"{ws.title}!{first}"
                chunks.append(
                    {
                        "loc_type": "sheet",
                        "locator": locator,
                        "text": f"Sheet {ws.title}: " + " | ".join(values),
                        "ordinal": ordinal,
                    }
                )
                ordinal += 1
        col_letter = get_column_letter(ws.max_column or 1)
        sheets.append({"name": ws.title, "rows": len(rows), "range": f"{ws.title}!A1:{col_letter}{ws.max_row or 1}"})
        if rows:
            chunks.append(
                {
                    "loc_type": "sheet",
                    "locator": f"{ws.title}!A1:{col_letter}{ws.max_row or 1}",
                    "text": f"Full sheet {ws.title}:\n" + "\n".join(rows),
                    "ordinal": ordinal,
                }
            )
            ordinal += 1
    # pandas preview for numeric tables
    try:
        xls = pd.ExcelFile(path)
        preview_parts = [f"Workbook {filename} sheets: {xls.sheet_names}"]
        for name in xls.sheet_names:
            df = pd.read_excel(path, sheet_name=name, header=None)
            preview_parts.append(f"\n## {name}\n{df.fillna('').to_string(index=False, header=False)}")
        preview = "\n".join(preview_parts)[:12000]
    except Exception:
        preview = "\n".join(c["text"] for c in chunks[:20])
    return {"kind": "xlsx", "sheets": sheets, "chunks": chunks, "preview": preview}


def _split_blocks(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []
    parts = re.split(r"\n\s*\n+", text)
    blocks = []
    buf = ""
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if len(buf) + len(part) < 900:
            buf = f"{buf}\n{part}".strip()
        else:
            if buf:
                blocks.append(buf)
            buf = part
    if buf:
        blocks.append(buf)
    return blocks or [text]


def _clean_pdf_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\uf0a0", " ")
    # collapse single-character-per-line artifacts
    lines = text.splitlines()
    rebuilt = []
    run = []
    for line in lines:
        stripped = line.strip()
        if len(stripped) <= 1:
            run.append(stripped)
            continue
        if run:
            rebuilt.append("".join(run))
            run = []
        rebuilt.append(stripped)
    if run:
        rebuilt.append("".join(run))
    text = "\n".join(rebuilt)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
